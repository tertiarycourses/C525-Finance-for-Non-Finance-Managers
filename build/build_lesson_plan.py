#!/usr/bin/env python3
"""Lesson Plan for C525 Finance for Non-Finance Managers (non-WSQ).

Converted from the WSQ/IBF parent TGS-2022602569 "Financial Analysis for
Non-Finance Managers". Topics follow the parent 1:1; the parent's 16 hrs
(incl. 2 hrs assessment) are rescaled proportionally onto the non-WSQ house
figure of 7.5 instructional hours per day, 9:30 am - 5:30 pm.
"""
import os, sys, datetime
sys.path.insert(0, os.path.expanduser("~/.claude/skills/non-wsq-courseware-build/build"))
import prodoc
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TITLE = "Finance for Non-Finance Managers"
CODE = "C525"
VERSION = "1.0"
EFFECTIVE = "21 September 2026"
TRAINER = "Dr Alfred Ang"

BRAND = RGBColor(0x1F, 0x6F, 0xEB); GREY = RGBColor(0x55, 0x5B, 0x66)
HEADER_FILL = "1F6FEB"; TOPIC_FILL = "E8F0FE"; LUNCH_FILL = "FFF4E5"; ADMIN_FILL = "F3F5F8"

OUTCOMES = [
    "LO1: Outline the process of financial statement analysis, including the balance sheet, "
    "income statement and cash flow statement, and the ratios derived from them.",
    "LO2: Assess an organisation's financial health from financial analysis, including "
    "profitability, liquidity, leverage, efficiency and cash flow.",
    "LO3: Evaluate an organisation's historical financial performance and benchmark it "
    "against industry, including for the financial services sector.",
    "LO4: Define indicators for investment suitability and apply capital budgeting methods "
    "to evaluate potential investment returns.",
]

# Topics follow the WSQ parent's outline 1:1 (same count, titles and order).
TOPICS = [
    ("01", "The different financial ratios", "LO1", "3-69"),
    ("02", "Profitability Analysis", "LO2", "70-81"),
    ("03", "Cash Flow Analysis", "LO2", "82-84"),
    ("04", "Projected financial and cash-flow statements", "LO2", "85-89"),
    ("05", "The process of financial statements analysis", "LO1", "96-110"),
    ("06", "Use financial analysis to determine the health of an organisation", "LO2", "90-95"),
    ("07", "Evaluate organisation's historical financial performance", "LO3", "111-119"),
    ("08", "Investment suitability based on financial analysis", "LO4", "120-123"),
    ("09", "Analysis for the financial services industry", "LO3", "124-127"),
    ("10", "Benchmarking of financial performance against industry", "LO3", "128-129"),
    ("11", "Use Capital Budgeting to evaluate potential investment returns", "LO4", "130-145"),
]

ACTIVITIES = [
    ("01", "Financial statement analysis and ratios", "Topic 1", "Day 1", "65 minutes"),
    ("02", "Cash flow reconstruction and liquidity", "Topic 3", "Day 1", "55 minutes"),
    ("03", "Projected income statement and cash flow", "Topic 4", "Day 1", "70 minutes"),
    ("04", "Trend analysis and industry benchmarking", "Topic 7, 10", "Day 2", "65 minutes"),
    ("05", "Capital budgeting and investment suitability", "Topic 11", "Day 2", "55 minutes"),
]

# (start, end, mins, kind, text). Teaching = 450 min/day; lunch 30 min; ends 17:30.
SCHEDULE = {
    1: ("Financial statements, ratios, profitability, cash flow and projections", [
        ("9:30", "9:45", 15, "admin", "Welcome, trainer and learner introductions, course outline and ground rules"),
        ("9:45", "11:30", 105, "topic", "Topic 1 - The different financial ratios: balance sheet, income statement, cash flow statement; liquidity, leverage, efficiency ratios"),
        ("11:30", "12:35", 65, "lab", "Activity 01 - Financial statement analysis and ratios (Northstar Supplies)"),
        ("12:35", "13:05", 30, "lunch", "Lunch break"),
        ("13:05", "14:05", 60, "topic", "Topic 2 - Profitability Analysis: gross, operating and net margin, ROA, ROE"),
        ("14:05", "14:40", 35, "topic", "Topic 3 - Cash Flow Analysis: improving cash flow management, monthly cash flow statements"),
        ("14:40", "15:35", 55, "lab", "Activity 02 - Cash flow reconstruction and liquidity (Harbour Services)"),
        ("15:35", "16:20", 45, "topic", "Topic 4 - Projected financial and cash-flow statements: proforma statements and forecasting methods"),
        ("16:20", "17:30", 70, "lab", "Activity 03 - Projected income statement and cash flow (Meridian Components)"),
    ]),
    2: ("Analysis process, financial health, historical performance, benchmarking and capital budgeting", [
        ("9:30", "9:45", 15, "admin", "Day 1 recap and Q&A"),
        ("9:45", "11:00", 75, "topic", "Topic 5 - The process of financial statements analysis: ratio, horizontal and vertical analysis"),
        ("11:00", "11:45", 45, "topic", "Topic 6 - Use financial analysis to determine the health of an organisation"),
        ("11:45", "12:30", 45, "topic", "Topic 7 - Evaluate organisation's historical financial performance: trend analysis case study"),
        ("12:30", "13:00", 30, "lunch", "Lunch break"),
        ("13:00", "13:35", 35, "topic", "Topic 8 - Investment suitability based on financial analysis"),
        ("13:35", "14:05", 30, "topic", "Topic 9 - Analysis for the financial services industry: banking and insurance ratios"),
        ("14:05", "14:30", 25, "topic", "Topic 10 - Benchmarking of financial performance against industry"),
        ("14:30", "15:35", 65, "lab", "Activity 04 - Trend analysis and industry benchmarking"),
        ("15:35", "16:20", 45, "topic", "Topic 11 - Use Capital Budgeting to evaluate potential investment returns: time value of money, payback, NPV, PI, IRR"),
        ("16:20", "17:15", 55, "lab", "Activity 05 - Capital budgeting and investment suitability"),
        ("17:15", "17:30", 15, "admin", "Course recap, Q&A and close"),
    ]),
}

doc = Document()
prodoc.style_headings(doc)
prodoc.add_cover_page(doc, "LESSON PLAN", TITLE, VERSION, course_code=CODE)
prodoc.add_version_control(doc, [
    (VERSION, EFFECTIVE,
     "Initial release of the non-WSQ lesson plan. Topics follow the parent course 1:1; "
     "schedule retimed to 9:30 am - 5:30 pm at 7.5 instructional hours per day.",
     TRAINER),
])
prodoc.add_toc(doc)

def H(text, level=1):
    p = doc.add_heading(text, level=level)
    for r in p.runs:
        r.font.name = "Arial"; r.font.color.rgb = BRAND
    return p

def set_cell(cell, text, bold=False, size=9.5, color=None, fill=None):
    cell.text = ""; p = cell.paragraphs[0]
    r = p.add_run(text); r.bold = bold; r.font.size = Pt(size); r.font.name = "Arial"
    if color: r.font.color.rgb = color
    if fill: prodoc._shade_cell(cell, fill)

H("Course Information", 1)
info = doc.add_table(rows=0, cols=2); info.style = "Table Grid"
for k, v in [
    ("Course Title", TITLE),
    ("Course Reference", CODE),
    ("Training Provider", "Tertiary Infotech Academy Pte Ltd  (UEN 201200696W)"),
    ("Duration", "2 days - 7.5 instructional hours per day (15 hours total)"),
    ("Daily Timing", "9:30 am - 5:30 pm (30-minute lunch break)"),
    ("Mode", "Instructor-led classroom or live online, with hands-on Excel activities"),
    ("Instructional Methods", "Lecture, demonstration, worked examples, guided practice and discussion"),
    ("Trainer", TRAINER),
]:
    c = info.add_row().cells
    set_cell(c[0], k, bold=True, size=10, fill=TOPIC_FILL); set_cell(c[1], v, size=10)
    c[0].width = Inches(1.9); c[1].width = Inches(4.9)

H("Learning Outcomes", 1)
doc.add_paragraph("On completion of this course, learners will be able to:")
for o in OUTCOMES:
    p = doc.add_paragraph(style="List Bullet"); p.add_run(o).font.size = Pt(10.5)

H("Course Schedule", 1)
p = doc.add_paragraph()
r = p.add_run("Each training day runs 9:30 am to 5:30 pm and totals 7.5 instructional hours "
              "(450 minutes) of teaching and hands-on activity, with a 30-minute lunch break.")
r.font.size = Pt(10.5)

for day, (theme, rows) in SCHEDULE.items():
    H(f"Day {day} - {theme}", 2)
    tbl = doc.add_table(rows=0, cols=3); tbl.style = "Table Grid"; tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = tbl.add_row().cells
    for i, h in enumerate(["Time", "Duration", "Topic / Activity"]):
        set_cell(hdr[i], h, bold=True, size=10, color=RGBColor(0xFF, 0xFF, 0xFF), fill=HEADER_FILL)
    teaching = 0
    for start, end, mins, kind, text in rows:
        fill = {"topic": TOPIC_FILL, "lunch": LUNCH_FILL, "admin": ADMIN_FILL}.get(kind)
        cells = tbl.add_row().cells
        set_cell(cells[0], f"{start}-{end}", bold=(kind == "topic"), fill=fill)
        set_cell(cells[1], f"{mins} min", fill=fill)
        set_cell(cells[2], text, bold=(kind == "topic"), fill=fill)
        if kind != "lunch": teaching += mins
    for row in tbl.rows:
        row.cells[0].width = Inches(1.05); row.cells[1].width = Inches(0.8); row.cells[2].width = Inches(4.95)
    pr = doc.add_paragraph().add_run(
        f"Total instructional time: {teaching} minutes ({teaching/60:g} hours), 9:30 am - 5:30 pm.")
    pr.italic = True; pr.font.size = Pt(9.5); pr.font.color.rgb = GREY
    assert teaching == 450, f"Day {day} instructional minutes = {teaching}, expected 450"

H("Topic and Activity Reference", 1)
tt = doc.add_table(rows=0, cols=4); tt.style = "Table Grid"
hdr = tt.add_row().cells
for i, h in enumerate(["Topic", "Outcome", "Activities", "Slides"]):
    set_cell(hdr[i], h, bold=True, size=10, color=RGBColor(0xFF, 0xFF, 0xFF), fill=HEADER_FILL)
for num, title, lo, slides in TOPICS:
    acts = [a[0] for a in ACTIVITIES if f"Topic {int(num)}" in a[2]]
    c = tt.add_row().cells
    set_cell(c[0], f"Topic {num}: {title}", bold=True, fill=TOPIC_FILL)
    set_cell(c[1], lo, fill=TOPIC_FILL)
    set_cell(c[2], ", ".join(f"Activity {a}" for a in acts) or "-")
    set_cell(c[3], slides)

H("Hands-On Activity Schedule", 1)
at = doc.add_table(rows=0, cols=5); at.style = "Table Grid"
hdr = at.add_row().cells
for i, h in enumerate(["Activity", "Title", "Topic", "Day", "Duration"]):
    set_cell(hdr[i], h, bold=True, size=10, color=RGBColor(0xFF, 0xFF, 0xFF), fill=HEADER_FILL)
for num, title, topic, day, dur in ACTIVITIES:
    c = at.add_row().cells
    set_cell(c[0], f"Activity {num}", bold=True, fill=TOPIC_FILL)
    set_cell(c[1], title); set_cell(c[2], topic); set_cell(c[3], day); set_cell(c[4], dur)

H("Learning Reinforcement", 1)
for a in ["Each topic is demonstrated with worked examples on real-shaped financial statements before learners practise.",
          "Each hands-on activity ends with a verification step so learners can confirm their own results.",
          "The trainer circulates during activity time to give individual feedback and unblock learners.",
          "Each topic closes with a recap mapped back to the course learning outcomes.",
          "Day 2 opens with a recap of Day 1 to consolidate the financial statement and ratio foundations."]:
    p = doc.add_paragraph(style="List Bullet"); p.add_run(a).font.size = Pt(10.5)

H("Resources Required", 1)
for r_ in ["A laptop with Microsoft Excel (or a compatible spreadsheet application).",
           "The activity workbooks and instruction sheets from the course GitHub repository.",
           "Trainer Slides, Learner Guide and Lesson Plan (downloadable from the LMS/TMS portal).",
           "Projector or screen-share for trainer demonstrations.",
           "Whiteboard or digital equivalent for worked calculations."]:
    p = doc.add_paragraph(style="List Bullet"); p.add_run(r_).font.size = Pt(10.5)

prodoc.add_page_numbers(doc)
prodoc.enable_update_fields(doc)
OUT = os.path.join(REPO, "courseware", f"LP-{TITLE} ({CODE}) v{VERSION}.docx")
doc.save(OUT)
print("Saved", OUT)
